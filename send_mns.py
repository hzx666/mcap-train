#!/usr/bin/env python3
"""
向 MNS 消费队列发送 mcap_extract_train 消息。

消息体与 run_extract_mns.py 读取的结构一致，默认对 body 做 Base64（与 Java SDK 相同）:

  {"event": "mcap_extract_train", "code": 200, "msg": "", "traceId": "", "data": {"ossPath": "oss://..."}}

用法:
  python send_mns.py --oss-path oss://bucket/a.mcap --trace-id t1
  python send_mns.py --oss-path oss://bucket/a.mcap --oss-path oss://bucket/b.mcap
  python send_mns.py --excel tasks.xlsx
  python send_mns.py --dry-run --oss-path oss://bucket/a.mcap --trace-id t1

环境变量（与 run_extract_mns.py 相同）:
  MNS_ENDPOINT            必填
  MNS_QUEUE_NAME          消费队列名（-k 优先级更高）
  MNS_ACCESS_KEY_ID       默认 OSS_ACCESS_KEY_ID
  MNS_ACCESS_KEY_SECRET   默认 OSS_ACCESS_KEY_SECRET
  MNS_SECURITY_TOKEN      可选 STS token
  MNS_EVENT               默认 mcap_extract_train
  MNS_BASE64              默认 true
"""

import argparse
import json
import os
import sys

MNS_ENDPOINT = os.environ.get("MNS_ENDPOINT", "").strip()
MNS_QUEUE_NAME = os.environ.get("MNS_QUEUE_NAME", "").strip()
MNS_ACCESS_KEY_ID = (
    os.environ.get("MNS_ACCESS_KEY_ID")
    or os.environ.get("OSS_ACCESS_KEY_ID")
    or ""
).strip()
MNS_ACCESS_KEY_SECRET = (
    os.environ.get("MNS_ACCESS_KEY_SECRET")
    or os.environ.get("OSS_ACCESS_KEY_SECRET")
    or ""
).strip()
MNS_SECURITY_TOKEN = os.environ.get("MNS_SECURITY_TOKEN", "").strip()
MNS_EVENT = os.environ.get("MNS_EVENT", "mcap_extract_train").strip() or "mcap_extract_train"

OSS_COL_ALIASES = {"osspath", "oss_path", "oss path"}
TRACE_COL_ALIASES = {"traceid", "trace_id", "trace id"}


def _env_bool(name: str, default: bool) -> bool:
    raw = os.environ.get(name)
    if raw is None or not str(raw).strip():
        return default
    return str(raw).strip().lower() in ("1", "true", "yes", "on")


MNS_BASE64 = _env_bool("MNS_BASE64", True)


def _norm_header(value) -> str:
    return str(value).strip().lower() if value is not None else ""


def build_message(oss_path: str, trace_id: str, event: str, code: int, msg: str) -> dict:
    data = {}
    if oss_path:
        data["ossPath"] = oss_path
    return {
        "event": event,
        "code": int(code),
        "msg": msg or "",
        "traceId": trace_id or "",
        "data": data,
    }


def read_excel_tasks(xlsx_path: str, sheet: str | None) -> list[tuple[str, str]]:
    """Return [(oss_path, trace_id), ...]. trace_id is empty when the column is absent."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        print("[ERROR] openpyxl 未安装，请执行: pip install openpyxl")
        sys.exit(1)

    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        ws = wb[sheet] if sheet else wb.active
        rows = ws.iter_rows(values_only=True)
        header_row = next(rows, None)
        if not header_row:
            print("[ERROR] Excel 为空")
            sys.exit(1)
        headers = [_norm_header(h) for h in header_row]
        oss_idx = next((i for i, h in enumerate(headers) if h in OSS_COL_ALIASES), None)
        if oss_idx is None:
            print(f"[ERROR] 找不到 oss_path / ossPath 列，现有列: {list(header_row)}")
            sys.exit(1)
        trace_idx = next((i for i, h in enumerate(headers) if h in TRACE_COL_ALIASES), None)

        tasks = []
        for row in rows:
            if oss_idx >= len(row) or row[oss_idx] is None:
                continue
            oss_path = str(row[oss_idx]).strip()
            if not oss_path:
                continue
            trace_id = ""
            if trace_idx is not None and trace_idx < len(row) and row[trace_idx] is not None:
                trace_id = str(row[trace_idx]).strip()
            tasks.append((oss_path, trace_id))
        return tasks
    finally:
        wb.close()


def open_queue(queue_name: str):
    try:
        from mns.account import Account
    except ImportError:
        print("[ERROR] mns 包未安装，请执行: pip install aliyun-mns-sdk")
        sys.exit(1)

    account = Account(
        MNS_ENDPOINT,
        MNS_ACCESS_KEY_ID,
        MNS_ACCESS_KEY_SECRET,
        MNS_SECURITY_TOKEN,
    )
    queue = account.get_queue(queue_name)
    queue.set_encoding(MNS_BASE64)
    return queue


def send_messages(queue, bodies: list[dict]) -> int:
    from mns.queue import Message

    sent = 0
    for body in bodies:
        resp = queue.send_message(Message(json.dumps(body, ensure_ascii=False)))
        sent += 1
        print(f"[INFO] sent message_id={resp.message_id} traceId={body.get('traceId')!r}")
    return sent


def main():
    parser = argparse.ArgumentParser(description="向 MNS 队列发送 mcap_extract_train 消息")
    parser.add_argument(
        "--oss-path", action="append", default=[],
        help="写入 data.ossPath，可重复传入以发送多条",
    )
    parser.add_argument("--trace-id", default="", help="写入 traceId；Excel 行里已有 traceId 时以表格为准")
    parser.add_argument("--event", default=MNS_EVENT, help="事件名（默认: mcap_extract_train）")
    parser.add_argument("--code", type=int, default=200, help="默认 200")
    parser.add_argument("--msg", default="", help="默认空字符串")
    parser.add_argument("--excel", default=None, help="从 xlsx 读取 ossPath，有 traceId 列则一并带上")
    parser.add_argument("--sheet", default=None, help="工作表名称（默认第一个 sheet）")
    parser.add_argument(
        "-k", "--queue", default=MNS_QUEUE_NAME,
        help="目标队列（默认环境变量 MNS_QUEUE_NAME）",
    )
    parser.add_argument("--dry-run", action="store_true", help="只打印消息体，不连接 MNS")
    args = parser.parse_args()

    tasks: list[tuple[str, str]] = []
    if args.excel:
        if not os.path.isfile(args.excel):
            print(f"[ERROR] 文件不存在: {args.excel}")
            sys.exit(1)
        tasks.extend(read_excel_tasks(args.excel, args.sheet))
    for oss_path in args.oss_path:
        tasks.append((oss_path.strip(), args.trace_id))

    if not tasks:
        print("[ERROR] 没有要发送的消息，请传入 --oss-path 或 --excel")
        sys.exit(1)

    bodies = []
    for oss_path, trace_id in tasks:
        if not trace_id:
            trace_id = args.trace_id
        bodies.append(build_message(oss_path, trace_id, args.event, args.code, args.msg))

    print(f"[INFO] 消息数   : {len(bodies)}")
    print(f"[INFO] Base64   : {MNS_BASE64}")
    if args.dry_run:
        for i, body in enumerate(bodies, 1):
            print(f"--- [{i}]")
            print(json.dumps(body, ensure_ascii=False))
        print("[INFO] Dry run: 未发送")
        return

    missing = []
    if not MNS_ENDPOINT:
        missing.append("MNS_ENDPOINT")
    if not args.queue:
        missing.append("MNS_QUEUE_NAME 或 --queue")
    if not MNS_ACCESS_KEY_ID:
        missing.append("MNS_ACCESS_KEY_ID 或 OSS_ACCESS_KEY_ID")
    if not MNS_ACCESS_KEY_SECRET:
        missing.append("MNS_ACCESS_KEY_SECRET 或 OSS_ACCESS_KEY_SECRET")
    if missing:
        print("[ERROR] 缺少配置: " + ", ".join(missing))
        sys.exit(1)

    print(f"[INFO] 队列     : {args.queue}")
    queue = open_queue(args.queue)
    sent = send_messages(queue, bodies)
    print(f"[INFO] 已发送   : {sent}")


if __name__ == "__main__":
    main()
